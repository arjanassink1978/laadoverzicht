import smtplib
from datetime import datetime, timedelta
from decimal import Decimal
from email import message_from_bytes, policy
from pathlib import Path
from unittest import mock

import pytest
from openpyxl import load_workbook

import laadoverzicht as lo

REAL_LOAD_CONFIG = lo.load_config


def make_cfg(tmp_path, **over):
    env = {
        "ZAPTEC_USERNAME": "u", "ZAPTEC_PASSWORD": "p",
        "GMAIL_USER": "me@gmail.com", "GMAIL_APP_PASSWORD": "pw",
        "MAIL_TO": "werkgever@example.com", "OWNER_NAME": "Jan", "OUTPUT_DIR": str(tmp_path),
    }
    env.update(over)
    return REAL_LOAD_CONFIG(env)


class FakeResp:
    def __init__(self, payload):
        self.payload = payload

    def raise_for_status(self):
        pass

    def json(self):
        return self.payload


def item(i, start, end, energy):
    return {"Id": f"s{i}", "StartDateTime": start, "EndDateTime": end, "Energy": energy,
            "DeviceName": "Go2"}


@pytest.fixture
def api():
    """Mock Zaptec: token + chargehistory met paginering; slaat de aanroepen op."""
    state = {"items": [], "calls": []}

    def post(url, **kw):
        assert url == lo.TOKEN_URL
        assert kw["data"]["grant_type"] == "password"
        return FakeResp({"access_token": "tok"})

    def get(url, headers=None, params=None, **kw):
        assert url == lo.HISTORY_URL
        assert headers["Authorization"] == "Bearer tok"
        state["calls"].append(params)
        size, idx = params["PageSize"], params["PageIndex"]
        return FakeResp({"Data": state["items"][idx * size:(idx + 1) * size]})

    with mock.patch.object(lo.requests, "post", post), mock.patch.object(lo.requests, "get", get):
        yield state


def test_previous_month():
    assert lo.previous_month(datetime(2026, 11, 1)) == (2026, 10)
    assert lo.previous_month(datetime(2026, 1, 1)) == (2025, 12)


def test_month_bounds_use_amsterdam_time():
    start, end = lo.month_bounds(2026, 10)
    assert start.isoformat() == "2026-09-30T22:00:00+00:00"  # zomertijd
    assert end.isoformat() == "2026-10-31T23:00:00+00:00"  # wintertijd na 25 okt
    start, end = lo.month_bounds(2026, 12)
    assert (start.isoformat(), end.isoformat()) == ("2026-11-30T23:00:00+00:00", "2026-12-31T23:00:00+00:00")


def test_fetch_sends_parameters_and_filters_month_edges(tmp_path, api):
    api["items"] = [
        item(1, "2026-09-30T21:59:59", "2026-09-30T22:30:00", 5),  # nog september (23:59 NL)
        item(2, "2026-09-30T22:00:00", "2026-09-30T23:00:00", 10),  # 1 okt 00:00 NL
        item(3, "2026-10-31T22:59:59Z", "2026-11-01T01:00:00Z", 7),  # 31 okt 23:59 NL
        item(4, "2026-10-31T23:00:00", "2026-11-01T01:00:00", 9),  # 1 nov 00:00 NL
    ]
    cfg = make_cfg(tmp_path, ZAPTEC_CHARGER_ID="abc")
    sessions = lo.fetch_sessions(cfg, 2026, 10)
    assert [s["kwh"] for s in sessions] == [Decimal(10), Decimal(7)]
    assert sessions[0]["start"].strftime("%d-%m %H:%M") == "01-10 00:00"
    p = api["calls"][0]
    assert p["From"] == "2026-09-30T22:00:00Z" and p["To"] == "2026-10-31T23:00:00Z"
    assert p["ChargerId"] == "abc"


def test_fetch_paginates(tmp_path, api):
    api["items"] = [item(i, "2026-10-05T10:00:00", "2026-10-05T11:00:00", 1) for i in range(lo.PAGE_SIZE + 5)]
    sessions = lo.fetch_sessions(make_cfg(tmp_path), 2026, 10)
    assert len(sessions) == lo.PAGE_SIZE + 5
    assert [c["PageIndex"] for c in api["calls"]] == [0, 1]
    assert "ChargerId" not in api["calls"][0]


def test_costs_and_total(tmp_path, api):
    api["items"] = [
        item(1, "2026-10-02T06:00:00", "2026-10-02T08:00:00", 12.5),   # 3.438 -> 3.44
        item(2, "2026-10-03T18:00:00", "2026-10-03T19:00:00", 10),     # 2.7505 -> 2.75
    ]
    cfg = make_cfg(tmp_path)
    path, kwh, total = lo.run(cfg, 2026, 10, dry_run=True)
    assert kwh == Decimal("22.5") and total == Decimal("6.19")
    ws = load_workbook(path).active
    rows = list(ws.iter_rows(min_row=7, values_only=True))
    assert rows[0][3] == 12.5 and rows[0][4] == 0.27505 and rows[0][5] == 3.44
    assert rows[1][5] == 2.75
    assert rows[2][0] == "Totaal" and rows[2][3] == 22.5 and rows[2][5] == 6.19
    assert ws["A1"].value == "Laadoverzicht oktober 2026" and ws["B2"].value == "Jan"
    assert ws["A7"].value.strftime("%d-%m-%Y") == "02-10-2026"
    assert ws["B7"].value.strftime("%H:%M") == "08:00"  # 06:00 UTC = 08:00 zomertijd


def test_custom_rate(tmp_path, api):
    api["items"] = [item(1, "2026-10-02T06:00:00", "2026-10-02T08:00:00", 10)]
    _, _, total = lo.run(make_cfg(tmp_path, RATE="0,30"), 2026, 10, dry_run=True)
    assert total == Decimal("3.00")


def test_cc_defaults_to_gmail_user(tmp_path):
    assert make_cfg(tmp_path)["mail_cc"] == "me@gmail.com"
    assert make_cfg(tmp_path, MAIL_CC="x@y.nl")["mail_cc"] == "x@y.nl"
    assert make_cfg(tmp_path, MAIL_CC="")["mail_cc"] == "me@gmail.com"


def test_defaults(tmp_path):
    cfg = make_cfg(tmp_path)
    assert cfg["owner"] == "Jan" and cfg["rate"] == Decimal("0.27505")


@pytest.mark.parametrize("dry_run", [False, True])
@pytest.mark.parametrize("value", ["", "  "])
def test_owner_name_required(tmp_path, value, dry_run):
    cfg = make_cfg(tmp_path, OWNER_NAME=value)
    with pytest.raises(lo.ConfigError, match="OWNER_NAME"):
        lo.validate(cfg, dry_run=dry_run)


def test_owner_name_missing():
    cfg = REAL_LOAD_CONFIG({"ZAPTEC_USERNAME": "u", "ZAPTEC_PASSWORD": "p"})
    with pytest.raises(lo.ConfigError, match="OWNER_NAME"):
        lo.validate(cfg, dry_run=True)


def test_mail_to_required_but_not_for_dry_run(tmp_path, api):
    cfg = make_cfg(tmp_path, MAIL_TO="")
    with pytest.raises(lo.ConfigError, match="MAIL_TO"):
        lo.validate(cfg, dry_run=False)
    lo.validate(cfg, dry_run=True)


def test_subject_and_filename(tmp_path):
    assert lo.subject_for("Jan", 2026, 10) == "declaratie Jan laadpaal oktober 2026"
    assert lo.subject_for("Piet", 2027, 3) == "declaratie Piet laadpaal maart 2027"


def run_with_smtp(cfg, api, dry_run=False):
    api["items"] = [item(1, "2026-10-02T06:00:00", "2026-10-02T08:00:00", 10)]
    with mock.patch.object(lo.smtplib, "SMTP") as smtp_cls:
        result = lo.run(cfg, 2026, 10, dry_run=dry_run)
    return result, smtp_cls


def test_mail_content(tmp_path, api):
    (path, _, _), smtp_cls = run_with_smtp(make_cfg(tmp_path, OWNER_NAME="Jan"), api)
    smtp_cls.assert_called_once_with("smtp.gmail.com", 587, timeout=lo.TIMEOUT)
    smtp = smtp_cls.return_value.__enter__.return_value
    smtp.starttls.assert_called_once()
    smtp.login.assert_called_once_with("me@gmail.com", "pw")
    sent = smtp.send_message.call_args
    msg = sent.args[0]
    assert sent.kwargs["to_addrs"] == ["werkgever@example.com", "me@gmail.com"]
    assert msg["Subject"] == "declaratie Jan laadpaal oktober 2026"
    assert msg["To"] == "werkgever@example.com" and msg["Cc"] == "me@gmail.com"
    body = msg.get_body().get_content()
    assert "oktober 2026" in body and "€ 2,75" in body
    assert "Totaal: 10,00 kWh" in body
    att = next(msg.iter_attachments())
    assert att.get_filename() == "declaratie Jan laadpaal oktober 2026.xlsx"
    assert path.name == att.get_filename() and path.parent == tmp_path
    assert att.get_content() == path.read_bytes()


def test_greeting_with_recipient_name(tmp_path, api):
    _, smtp_cls = run_with_smtp(make_cfg(tmp_path, MAIL_TO_NAME="Wim"), api)
    msg = smtp_cls.return_value.__enter__.return_value.send_message.call_args.args[0]
    assert msg.get_body().get_content().startswith("Beste Wim,\n")


def test_greeting_without_recipient_name(tmp_path, api):
    _, smtp_cls = run_with_smtp(make_cfg(tmp_path, MAIL_TO_NAME=""), api)
    msg = smtp_cls.return_value.__enter__.return_value.send_message.call_args.args[0]
    assert msg.get_body().get_content().startswith("Beste,\n")


def test_attachment_filename_survives_serialisation(tmp_path, api):
    """Spaties in de bijlagenaam moeten na verzenden (RFC 2231/quoted) heel blijven."""
    run_with_smtp(make_cfg(tmp_path), api)
    with mock.patch.object(lo.smtplib, "SMTP") as smtp_cls:
        lo.run(make_cfg(tmp_path), 2026, 10)
    msg = smtp_cls.return_value.__enter__.return_value.send_message.call_args.args[0]
    parsed = message_from_bytes(msg.as_bytes(), policy=policy.default)
    att = next(parsed.iter_attachments())
    assert att.get_filename() == "declaratie Jan laadpaal oktober 2026.xlsx"
    assert att.get_content_type() == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_dry_run_builds_xlsx_without_mail(tmp_path, api):
    (path, _, _), smtp_cls = run_with_smtp(make_cfg(tmp_path), api, dry_run=True)
    assert path.exists()
    smtp_cls.assert_not_called()


def test_xlsx_date_and_time_formats_are_lowercase(tmp_path, api):
    api["items"] = [item(1, "2026-10-02T06:00:00", "2026-10-02T08:00:00", 10)]
    (path, _, _), _ = run_with_smtp(make_cfg(tmp_path), api, dry_run=True)
    ws = load_workbook(path).active
    assert ws["A7"].number_format == "dd-mm-yyyy"
    assert ws["B7"].number_format == "hh:mm"
    assert ws["C7"].number_format == "hh:mm"


def test_main_dry_run_cli(tmp_path, api, monkeypatch):
    api["items"] = [item(1, "2026-10-02T06:00:00", "2026-10-02T08:00:00", 10)]
    monkeypatch.setattr(lo, "load_config", lambda: make_cfg(tmp_path))
    with mock.patch.object(lo.smtplib, "SMTP") as smtp_cls:
        assert lo.main(["--month", "2026-10", "--dry-run"]) == 0
    smtp_cls.assert_not_called()
    assert (tmp_path / "declaratie Jan laadpaal oktober 2026.xlsx").exists()
    assert (tmp_path / "laadoverzicht.log").exists()


def test_main_reports_config_error(tmp_path, monkeypatch):
    monkeypatch.setattr(lo, "load_config", lambda: make_cfg(tmp_path, MAIL_TO=""))
    assert lo.main(["--month", "2026-10"]) == 2


def test_parse_month_rejects_garbage():
    with pytest.raises(Exception):
        lo.parse_month("2026-13")


# --- dynamisch tarief ---

def price_payload(first_hour, prices):
    """EnergyZero-antwoord: uurprijzen vanaf first_hour (UTC) oplopend."""
    t = datetime.fromisoformat(first_hour)
    return [{"readingDate": (t + timedelta(hours=i)).strftime("%Y-%m-%dT%H:%M:%SZ"), "price": p}
            for i, p in enumerate(prices)]


@pytest.fixture
def dyn(api):
    """Zaptec-mock uit `api` plus gemockte EnergyZero; zet de prijzen in dyn['prices']."""
    state = {"prices": [], "price_calls": [], "zaptec": api}
    zaptec_get = lo.requests.get

    def get(url, **kw):
        if url == lo.PRICES_URL:
            state["price_calls"].append(kw["params"])
            return FakeResp({"Prices": state["prices"]})
        return zaptec_get(url, **kw)

    with mock.patch.object(lo.requests, "get", get):
        yield state


def dyn_cfg(tmp_path, **over):
    env = {"RATE_MODE": "dynamic", "ENERGY_TAX": "0.10", "VAT": "0.21", "DYNAMIC_MARKUP": "0.02"}
    env.update(over)
    return make_cfg(tmp_path, **env)


def dyn_item(i, start, end, energy, details):
    d = item(i, start, end, energy)
    d["EnergyDetails"] = [{"Timestamp": t, "Energy": e} for t, e in details]
    return d


def test_fixed_mode_default_unchanged(tmp_path, api):
    cfg = make_cfg(tmp_path)
    assert cfg["rate_mode"] == "fixed"
    assert (cfg["dynamic_markup"], cfg["energy_tax"], cfg["vat"]) == (Decimal(0), Decimal("0.09161"), Decimal("0.21"))
    api["items"] = [item(1, "2026-10-02T06:00:00", "2026-10-02T08:00:00", 10)]
    lo.run(cfg, 2026, 10, dry_run=True)  # de api-mock faalt bij elke andere URL dan Zaptec, dus geen prijs-API
    assert "DetailLevel" not in api["calls"][0]
    ws = load_workbook(next(tmp_path.glob("*.xlsx"))).active
    assert ws["E6"].value == "Tarief (EUR/kWh)" and ws["E7"].value == 0.27505


def test_invalid_rate_mode_and_numbers(tmp_path):
    for k, v in (("RATE_MODE", "variabel"), ("VAT", "x"), ("ENERGY_TAX", "-1"), ("DYNAMIC_MARKUP", "abc")):
        with pytest.raises(lo.ConfigError, match=k):
            make_cfg(tmp_path, **{k: v})


def test_dynamic_rate_formula(tmp_path):
    # (0.08 + 0.02 + 0.10) * 1.21
    assert lo.dynamic_rate(Decimal("0.08"), dyn_cfg(tmp_path)) == Decimal("0.2420")


def test_dynamic_per_interval(tmp_path, dyn):
    # 10:00-12:00 UTC; 2 kWh in uur 10, 4 kWh in uur 11 (cumulatief 0, 2, 6)
    dyn["zaptec"]["items"] = [dyn_item(1, "2026-10-02T10:00:00", "2026-10-02T12:00:00", 6, [
        ("2026-10-02T10:00:00", 0), ("2026-10-02T11:00:00", 2), ("2026-10-02T12:00:00", 6)])]
    dyn["prices"] = price_payload("2026-10-02T09:00:00", [0.50, 0.10, 0.20])
    path, kwh, total = lo.run(dyn_cfg(tmp_path), 2026, 10, dry_run=True)
    # 2*(0.10+0.02+0.10)*1.21 + 4*(0.20+0.02+0.10)*1.21 = 0.5324 + 1.5488 = 2.0812
    assert kwh == Decimal(6) and total == Decimal("2.08")
    assert dyn["zaptec"]["calls"][0]["DetailLevel"] == 1
    assert dyn["price_calls"][0]["interval"] == 4 and dyn["price_calls"][0]["inclBtw"] == "false"
    ws = load_workbook(path).active
    assert ws["E6"].value == "Dynamisch tarief (gem. EUR/kWh)"
    assert ws["F7"].value == 2.08 and ws["E7"].value == pytest.approx(2.08 / 6)
    assert ws["A8"].value == "Totaal" and ws["F8"].value == 2.08


def test_dynamic_interval_spanning_hours_split_by_time(tmp_path, dyn):
    # één interval 10:30-11:30, 2 kWh: halve kWh per uur
    dyn["zaptec"]["items"] = [dyn_item(1, "2026-10-02T10:30:00", "2026-10-02T11:30:00", 2, [
        ("2026-10-02T10:30:00", 0), ("2026-10-02T11:30:00", 2)])]
    dyn["prices"] = price_payload("2026-10-02T10:00:00", [0.10, 0.30])
    _, _, total = lo.run(dyn_cfg(tmp_path, DYNAMIC_MARKUP="0", ENERGY_TAX="0", VAT="0"), 2026, 10, dry_run=True)
    assert total == Decimal("0.40")  # 1*0.10 + 1*0.30


def test_dynamic_fallback_even_distribution_warns(tmp_path, dyn, caplog):
    dyn["zaptec"]["items"] = [item(1, "2026-10-02T10:00:00", "2026-10-02T12:00:00", 4)]  # geen EnergyDetails
    dyn["prices"] = price_payload("2026-10-02T10:00:00", [0.10, 0.30])
    with caplog.at_level("WARNING"):
        _, _, total = lo.run(dyn_cfg(tmp_path, DYNAMIC_MARKUP="0", ENERGY_TAX="0", VAT="0"), 2026, 10, dry_run=True)
    assert total == Decimal("0.80")  # 2*0.10 + 2*0.30
    assert "gelijkmatig verdeeld" in caplog.text


def test_dynamic_empty_details_also_falls_back(tmp_path, dyn, caplog):
    dyn["zaptec"]["items"] = [dyn_item(1, "2026-10-02T10:00:00", "2026-10-02T11:00:00", 3, [])]
    dyn["prices"] = price_payload("2026-10-02T10:00:00", [0.10])
    with caplog.at_level("WARNING"):
        _, _, total = lo.run(dyn_cfg(tmp_path, DYNAMIC_MARKUP="0", ENERGY_TAX="0", VAT="0"), 2026, 10, dry_run=True)
    assert total == Decimal("0.30") and "gelijkmatig verdeeld" in caplog.text


def test_dynamic_missing_price_fails(tmp_path, dyn):
    dyn["zaptec"]["items"] = [item(1, "2026-10-02T10:00:00", "2026-10-02T12:00:00", 4)]
    dyn["prices"] = price_payload("2026-10-02T10:00:00", [0.10])  # uur 11 ontbreekt
    with pytest.raises(lo.PriceError, match="geen marktprijs"):
        lo.run(dyn_cfg(tmp_path), 2026, 10, dry_run=True)
    assert not list(tmp_path.glob("*.xlsx"))


def test_dynamic_rejects_sub_hour_price_resolution(tmp_path, dyn):
    dyn["zaptec"]["items"] = [item(1, "2026-10-02T10:00:00", "2026-10-02T11:00:00", 1)]
    dyn["prices"] = [{"readingDate": "2026-10-02T10:15:00Z", "price": 0.1}]
    with pytest.raises(lo.PriceError, match="uurprijzen"):
        lo.run(dyn_cfg(tmp_path), 2026, 10, dry_run=True)


def test_dynamic_vat_and_energy_tax(tmp_path, dyn):
    dyn["zaptec"]["items"] = [item(1, "2026-10-02T10:00:00", "2026-10-02T11:00:00", 10)]
    dyn["prices"] = price_payload("2026-10-02T10:00:00", [0.08])
    # (0.08 + 0.0049 + 0.09161) * 1.21 = 0.2135 per kWh -> 2.135 -> afgerond 2.14
    cfg = dyn_cfg(tmp_path, DYNAMIC_MARKUP="0.0049", ENERGY_TAX="0.09161")
    _, _, total = lo.run(cfg, 2026, 10, dry_run=True)
    assert total == Decimal("2.14")
    # zonder btw
    _, _, total = lo.run(dyn_cfg(tmp_path, DYNAMIC_MARKUP="0", ENERGY_TAX="0.10", VAT="0"), 2026, 10, dry_run=True)
    assert total == Decimal("1.80")


def test_number_nl():
    assert lo.number_nl(Decimal("176.19")) == "176,19"
    assert lo.number_nl(Decimal("1234.5")) == "1.234,50"
    assert lo.number_nl(Decimal("2"), 3) == "2,000"


def test_mail_kwh_decimal_comma(tmp_path):
    xlsx = tmp_path / "x.xlsx"
    xlsx.write_bytes(b"x")
    with mock.patch.object(lo.smtplib, "SMTP") as smtp_cls:
        lo.send_mail(make_cfg(tmp_path), 2026, 10, xlsx, Decimal("48.46"), Decimal("176.19"), 3)
    msg = smtp_cls.return_value.__enter__.return_value.send_message.call_args.args[0]
    body = msg.get_body().get_content()
    assert "Totaal: 176,19 kWh" in body and "€ 48,46" in body
