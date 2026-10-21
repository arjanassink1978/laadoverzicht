#!/usr/bin/env python3
"""Maandelijks Zaptec-laadoverzicht (xlsx) maken en mailen naar de werkgever."""
import argparse
import logging
import os
import smtplib
import sys
from datetime import date, datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from email.message import EmailMessage
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

BASE_DIR = Path(__file__).resolve().parent
TOKEN_URL = "https://api.zaptec.com/oauth/token"
HISTORY_URL = "https://api.zaptec.com/api/chargehistory"
SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 587
TZ = ZoneInfo("Europe/Amsterdam")
PAGE_SIZE = 100
TIMEOUT = 30
MONTHS_NL = [
    "januari", "februari", "maart", "april", "mei", "juni", "juli",
    "augustus", "september", "oktober", "november", "december",
]
XLSX_MIME = ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")

log = logging.getLogger("laadoverzicht")


class ConfigError(Exception):
    pass


def previous_month(today=None):
    today = today or date.today()
    return (today.year - 1, 12) if today.month == 1 else (today.year, today.month - 1)


def parse_month(value):
    try:
        d = datetime.strptime(value, "%Y-%m")
    except ValueError:
        raise argparse.ArgumentTypeError(f"ongeldige maand '{value}', gebruik YYYY-MM")
    return d.year, d.month


def month_label(year, month):
    return f"{MONTHS_NL[month - 1]} {year}"


def month_bounds(year, month):
    """Begin (inclusief) en einde (exclusief) van de kalendermaand in Europe/Amsterdam, als UTC."""
    start = datetime(year, month, 1, tzinfo=TZ)
    end = datetime(year + (month == 12), month % 12 + 1, 1, tzinfo=TZ)
    return start.astimezone(timezone.utc), end.astimezone(timezone.utc)


def subject_for(owner, year, month):
    return f"declaratie {owner} laadpaal {month_label(year, month)}"


def load_config(env=None):
    if env is None:
        load_dotenv(BASE_DIR / ".env")
        env = os.environ
    get = lambda k, d="": (env.get(k) or d).strip()
    cfg = {
        "zaptec_username": get("ZAPTEC_USERNAME"),
        "zaptec_password": get("ZAPTEC_PASSWORD"),
        "charger_id": get("ZAPTEC_CHARGER_ID"),
        "gmail_user": get("GMAIL_USER"),
        "gmail_password": get("GMAIL_APP_PASSWORD"),
        "mail_to": get("MAIL_TO"),
        "mail_to_name": get("MAIL_TO_NAME"),
        "owner": get("OWNER_NAME"),
        "output_dir": Path(get("OUTPUT_DIR", str(BASE_DIR / "output"))),
    }
    cfg["mail_cc"] = get("MAIL_CC") or cfg["gmail_user"]
    try:
        cfg["rate"] = Decimal(get("RATE", "0.27505").replace(",", "."))
    except Exception:
        raise ConfigError(f"RATE is geen geldig getal: {get('RATE')!r}")
    if cfg["rate"] < 0:
        raise ConfigError("RATE mag niet negatief zijn")
    cfg["log_file"] = Path(get("LOG_FILE")) if get("LOG_FILE") else cfg["output_dir"] / "laadoverzicht.log"
    return cfg


def validate(cfg, dry_run):
    required = ["zaptec_username", "zaptec_password", "owner"]
    if not dry_run:
        required += ["gmail_user", "gmail_password", "mail_to"]
    names = {"zaptec_username": "ZAPTEC_USERNAME", "zaptec_password": "ZAPTEC_PASSWORD", "owner": "OWNER_NAME",
             "gmail_user": "GMAIL_USER", "gmail_password": "GMAIL_APP_PASSWORD", "mail_to": "MAIL_TO"}
    missing = [names[k] for k in required if not cfg[k]]
    if missing:
        raise ConfigError("ontbrekende instellingen: " + ", ".join(missing))


def get_token(cfg):
    r = requests.post(
        TOKEN_URL,
        data={"grant_type": "password", "username": cfg["zaptec_username"],
              "password": cfg["zaptec_password"]},
        timeout=TIMEOUT,
    )
    r.raise_for_status()
    return r.json()["access_token"]


def parse_dt(value):
    """Zaptec geeft UTC-tijden, soms zonder offset; naive waarden zijn UTC."""
    dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt


def fetch_sessions(cfg, year, month, token=None):
    """Alle laadsessies die in de maand zijn gestart, oplopend gesorteerd."""
    start, end = month_bounds(year, month)
    token = token or get_token(cfg)
    headers = {"Authorization": f"Bearer {token}", "Accept": "application/json"}
    fmt = "%Y-%m-%dT%H:%M:%SZ"
    params = {"From": start.strftime(fmt), "To": end.strftime(fmt), "PageSize": PAGE_SIZE,
              "SortDescending": "false"}
    if cfg["charger_id"]:
        params["ChargerId"] = cfg["charger_id"]
    sessions, seen, page = [], set(), 0
    while True:
        r = requests.get(HISTORY_URL, headers=headers, timeout=TIMEOUT,
                         params={**params, "PageIndex": page})
        r.raise_for_status()
        data = r.json().get("Data") or []
        for item in data:
            if item.get("Id") in seen:
                continue
            seen.add(item.get("Id"))
            s = parse_dt(item["StartDateTime"])
            if not start <= s < end:
                continue
            e = parse_dt(item["EndDateTime"]) if item.get("EndDateTime") else s
            sessions.append({
                "start": s.astimezone(TZ),
                "end": e.astimezone(TZ),
                "kwh": Decimal(str(item.get("Energy") or 0)),
                "charger": item.get("DeviceName") or item.get("ChargerId") or "",
            })
        if len(data) < PAGE_SIZE:
            break
        page += 1
    sessions.sort(key=lambda s: s["start"])
    return sessions


def money(value):
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def priced(sessions, rate):
    """Kosten per sessie op centen afgerond; totaal is de som van de getoonde regels."""
    rows = [dict(s, cost=money(s["kwh"] * rate)) for s in sessions]
    return rows, sum((r["kwh"] for r in rows), Decimal(0)), sum((r["cost"] for r in rows), Decimal(0))


def build_xlsx(cfg, year, month, sessions, path):
    rows, total_kwh, total_cost = priced(sessions, cfg["rate"])
    charger = cfg["charger_id"] or ", ".join(sorted({s["charger"] for s in rows if s["charger"]})) or "Zaptec Go2"
    wb = Workbook()
    ws = wb.active
    ws.title = "Laadsessies"
    bold = Font(bold=True)
    ws["A1"] = f"Laadoverzicht {month_label(year, month)}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"], ws["B2"] = "Naam", cfg["owner"]
    ws["A3"], ws["B3"] = "Maand", month_label(year, month)
    ws["A4"], ws["B4"] = "Laadpaal", charger
    for r in (2, 3, 4):
        ws[f"A{r}"].font = bold
    header_row = 6
    for col, name in enumerate(["Datum", "Starttijd", "Eindtijd", "kWh", "Tarief (EUR/kWh)", "Kosten"], 1):
        c = ws.cell(header_row, col, name)
        c.font = bold
        c.alignment = Alignment(horizontal="center")
    r = header_row
    for s in rows:
        r += 1
        ws.cell(r, 1, s["start"].date()).number_format = "dd-mm-yyyy"
        ws.cell(r, 2, s["start"].time().replace(microsecond=0, tzinfo=None)).number_format = "hh:mm"
        ws.cell(r, 3, s["end"].time().replace(microsecond=0, tzinfo=None)).number_format = "hh:mm"
        ws.cell(r, 4, float(s["kwh"])).number_format = "0.00"
        ws.cell(r, 5, float(cfg["rate"])).number_format = '"€" 0.00000'
        ws.cell(r, 6, float(s["cost"])).number_format = '"€" #,##0.00'
    r += 1
    ws.cell(r, 1, "Totaal").font = bold
    kwh = ws.cell(r, 4, float(total_kwh))
    kwh.number_format, kwh.font = "0.00", bold
    tot = ws.cell(r, 6, float(total_cost))
    tot.number_format, tot.font = '"€" #,##0.00', bold
    for col, width in zip("ABCDEF", (14, 11, 11, 10, 18, 12)):
        ws.column_dimensions[col].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return total_kwh, total_cost


def euro_nl(value):
    s = f"{value:,.2f}"
    return "€ " + s.replace(",", "X").replace(".", ",").replace("X", ".")


def send_mail(cfg, year, month, xlsx_path, total_cost, total_kwh, count):
    subject = subject_for(cfg["owner"], year, month)
    msg = EmailMessage()
    msg["From"] = cfg["gmail_user"]
    msg["To"] = cfg["mail_to"]
    msg["Cc"] = cfg["mail_cc"]
    msg["Subject"] = subject
    greeting = f"Beste {cfg['mail_to_name']}," if cfg.get("mail_to_name") else "Beste,"
    msg.set_content(
        f"{greeting}\n\nIn de bijlage het laadoverzicht van {month_label(year, month)}.\n\n"
        f"Maand: {month_label(year, month)}\n"
        f"Aantal laadsessies: {count}\n"
        f"Totaal: {total_kwh:.2f} kWh\n"
        f"Totaalbedrag: {euro_nl(total_cost)}\n\n"
        f"Met vriendelijke groet,\n{cfg['owner']}\n"
    )
    msg.add_attachment(xlsx_path.read_bytes(), maintype=XLSX_MIME[0], subtype=XLSX_MIME[1],
                       filename=xlsx_path.name)
    recipients = [a.strip() for a in f"{cfg['mail_to']},{cfg['mail_cc']}".split(",") if a.strip()]
    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=TIMEOUT) as smtp:
        smtp.starttls()
        smtp.login(cfg["gmail_user"], cfg["gmail_password"])
        smtp.send_message(msg, to_addrs=recipients)


def run(cfg, year, month, dry_run=False, token=None):
    validate(cfg, dry_run)
    sessions = fetch_sessions(cfg, year, month, token)
    log.info("%d laadsessies in %s", len(sessions), month_label(year, month))
    path = cfg["output_dir"] / (subject_for(cfg["owner"], year, month) + ".xlsx")
    total_kwh, total_cost = build_xlsx(cfg, year, month, sessions, path)
    log.info("xlsx geschreven: %s (totaal %s)", path, euro_nl(total_cost))
    if dry_run:
        log.info("dry-run: geen mail verstuurd")
    else:
        send_mail(cfg, year, month, path, total_cost, total_kwh, len(sessions))
        log.info("mail verstuurd naar %s (cc %s)", cfg["mail_to"], cfg["mail_cc"])
    return path, total_kwh, total_cost


def setup_logging(log_file):
    handlers = [logging.StreamHandler()]
    try:
        log_file.parent.mkdir(parents=True, exist_ok=True)
        handlers.append(logging.FileHandler(log_file, encoding="utf-8"))
    except OSError as exc:
        print(f"kan logbestand {log_file} niet openen: {exc}", file=sys.stderr)
    logging.basicConfig(level=logging.INFO, handlers=handlers, force=True,
                        format="%(asctime)s %(levelname)s %(message)s")


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--month", type=parse_month, help="maand als YYYY-MM (standaard: vorige maand)")
    p.add_argument("--dry-run", action="store_true", help="maak de xlsx maar verstuur geen mail")
    args = p.parse_args(argv)
    try:
        cfg = load_config()
        setup_logging(cfg["log_file"])
        year, month = args.month or previous_month()
        run(cfg, year, month, args.dry_run)
    except ConfigError as exc:
        print(f"configuratiefout: {exc}", file=sys.stderr)
        return 2
    except Exception:
        log.exception("laadoverzicht mislukt")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
